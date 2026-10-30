# -*- coding: utf-8 -*-
import openpyxl, warnings, json, glob
warnings.filterwarnings("ignore")
def load(p, s, idk="缺陷ID"):
    wb = openpyxl.load_workbook(p); ws = wb[s]
    h = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    out = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d = dict(zip(h, r))
        if d.get(idk): out.append(d)
    wb.close(); return out
def mk(rows, idk, tk, dk, skip_sec=True, tm="测试模块", om="所属模块", env="测试环境", ct="创建时间"):
    fx = []
    for d in rows:
        g = lambda k: str(d.get(k) or "").strip()
        if skip_sec and g(tk).startswith("[安全]"): continue
        fx.append({"defectId": g(idk), "title": g(tk), "description": g(dk),
                   "status": g("缺陷状态"), "severity": g("严重级别"), "quality": g("质量特性"),
                   "testModule": g(tm), "moduleTag": g(om),
                   "duplicateOf": g("关联重复缺陷编号"), "createdAt": g(ct)[:10], "reproEnv": g(env)})
    return fx
out = {}
out["hcp"] = mk(load("缺陷列表demo/hcp-bug-1029/hcp.xlsx", "缺陷列表"), "缺陷ID", "缺陷标题", "缺陷描述")
hc = []
for i in range(1, 7): hc += load(f"缺陷列表demo/hc-bug-1029/hc{i}.xlsx", "缺陷列表")
out["hc"] = mk(hc, "缺陷ID", "缺陷标题", "缺陷描述")
hit = load("缺陷列表demo/hitools/hitools.xlsx", "缺陷列表")
out["hitools"] = mk(hit, "缺陷ID", "缺陷标题", "缺陷描述")
pe = {}
for p in sorted(glob.glob("缺陷列表demo/配套-交通项目*.xlsx")):
    for d in load(p, "Sheet1"): pe[str(d["缺陷ID"])] = d
out["pe"] = mk(list(pe.values()), "缺陷ID", "标题", "缺陷描述", skip_sec=False, ct="申请时间")
led = load("缺陷列表demo/led-客户端.xlsx", "缺陷列表")
out["led"] = mk(led, "缺陷ID", "缺陷标题", "缺陷描述", skip_sec=False)
for k, v in out.items():
    json.dump({"defects": v}, open(f".run1029/{k}.json", "w", encoding="utf-8"), ensure_ascii=False)
    print(f"{k:8} {len(v)}")
